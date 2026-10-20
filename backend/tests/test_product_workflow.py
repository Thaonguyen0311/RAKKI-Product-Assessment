import io
import importlib.util
import json
from types import SimpleNamespace
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from PIL import Image
import sqlalchemy as sa
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
from alembic.migration import MigrationContext
from alembic.operations import Operations

from app.db.session import Base, get_db
from app.main import app
from app.models.product import AIExtraction, Assessment, EvidenceSource, Material, Product, ProductImage, Revision
from app.core.config import settings


@pytest.fixture
def client(tmp_path, monkeypatch):
    engine = create_engine('sqlite://', connect_args={'check_same_thread': False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)

    def override_db():
        session = session_factory()
        try:
            yield session
        finally:
            session.close()

    monkeypatch.setattr(settings, 'storage_dir', str(tmp_path / 'images'))
    app.dependency_overrides[get_db] = override_db
    with TestClient(app) as test_client:
        yield test_client, session_factory, tmp_path / 'images'
    app.dependency_overrides.clear()
    Base.metadata.drop_all(engine)
    engine.dispose()


def png_bytes(color='red'):
    output = io.BytesIO()
    Image.new('RGB', (12, 12), color).save(output, format='PNG')
    return output.getvalue()


def test_canonical_material_mapping_is_complete_and_drives_product_fields(client):
    test_client, _, _ = client
    materials = test_client.get('/api/materials')
    assert materials.status_code == 200
    mapping = materials.json()
    assert len(mapping) == 304
    assert len({item['material_id'] for item in mapping}) == 304
    assert next(item for item in mapping if item['material'] == 'Crushed concrete') == {
        'material': 'Crushed concrete', 'material_id': 7000000187,
        'material_category': 'Infra, yard, and foundations | Concrete',
    }

    created = test_client.post('/api/products', json={'name': 'Aggregate', 'material': '  Crushed   concrete  ', 'material_id': 123})
    assert created.status_code == 201
    product = created.json()
    assert (product['material'], product['material_id'], product['material_category']) == (
        'Crushed concrete', 7000000187, 'Infra, yard, and foundations | Concrete')
    assert product['material_verification'] == 'verified'
    assert test_client.patch(f"/api/products/{product['id']}", json={'material': 'Fly ash'}).json()['material_id'] == 7000000184
    assert test_client.patch(f"/api/products/{product['id']}", json={'material': 'made-up material'}).status_code == 422
    reloaded = test_client.get(f"/api/products/{product['id']}").json()['product']
    assert (reloaded['material'], reloaded['material_id'], reloaded['material_category']) == (
        'Fly ash', 7000000184, 'Infra, yard, and foundations')
    export = test_client.get('/api/export.xlsx')
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(export.content), read_only=True).worksheets[0]
    assert 'Material ID' in [cell.value for cell in next(sheet.iter_rows(max_row=1))]


def test_analysis_resolves_material_suggestion_from_canonical_mapping(client, monkeypatch):
    test_client, _, _ = client
    product = test_client.post('/api/products', json={'name': ''}).json()
    uploaded = test_client.post(f"/api/products/{product['id']}/images", files={'file': ('item.png', png_bytes(), 'image/png')}).json()
    from app.services import gemini
    monkeypatch.setattr(gemini, 'analyze_image', lambda _path, _context=None: {
        'product_name': 'Crushed concrete sample', 'category': 'Building Components and Materials',
        'additional_notes': None, 'overall_condition': None, 'is_intact': None, 'is_clean': None,
        'confidence': 0.84, 'suggested_material': ' Crushed   concrete ', 'material_confidence': 0.77,
    })
    result = test_client.post(f"/api/products/{product['id']}/analyze?image_id={uploaded['uploaded'][0]['id']}")
    assert result.status_code == 200
    suggestion = result.json()['data']['material_suggestion']
    assert suggestion == {'suggested_name': 'Crushed concrete', 'material': 'Crushed concrete',
                          'material_id': 7000000187, 'material_category': 'Infra, yard, and foundations | Concrete',
                          'confidence': 0.77, 'status': 'suggested'}


def test_primary_image_migration_preserves_and_marks_existing_rows():
    engine = create_engine('sqlite://')
    with engine.begin() as connection:
        connection.exec_driver_sql('CREATE TABLE product_images (id INTEGER PRIMARY KEY, product_id INTEGER NOT NULL)')
        connection.exec_driver_sql('INSERT INTO product_images (id, product_id) VALUES (1, 7), (2, 7), (3, 8)')
        migration_path = Path(__file__).parents[1] / 'alembic' / 'versions' / '0002_primary_product_image.py'
        spec = importlib.util.spec_from_file_location('primary_image_migration', migration_path)
        migration = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(migration)
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()
        rows = connection.exec_driver_sql('SELECT id, product_id, is_primary FROM product_images ORDER BY id').all()
    assert rows == [(1, 7, 1), (2, 7, 0), (3, 8, 1)]
    engine.dispose()


def test_provenance_migration_keeps_legacy_product_values():
    engine = create_engine('sqlite://')
    with engine.begin() as connection:
        connection.exec_driver_sql('CREATE TABLE products (id INTEGER PRIMARY KEY, name TEXT, category TEXT, subtype TEXT, brand TEXT, sku TEXT, color TEXT, description TEXT, condition TEXT)')
        connection.exec_driver_sql("INSERT INTO products (id, name, brand) VALUES (1, 'AI jacket', 'Known brand'), (2, 'Legacy item', NULL)")
        connection.exec_driver_sql('CREATE TABLE ai_extractions (id INTEGER PRIMARY KEY, product_id INTEGER, payload JSON, created_at TEXT)')
        connection.exec_driver_sql("INSERT INTO ai_extractions (id, product_id, payload, created_at) VALUES (1, 1, '{\"name\":\"AI jacket\",\"brand\":null}', '2026-01-01')")
        migration_path = Path(__file__).parents[1] / 'alembic' / 'versions' / '0003_product_field_provenance.py'
        spec = importlib.util.spec_from_file_location('field_provenance_migration', migration_path)
        migration = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(migration)
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()
        rows = connection.exec_driver_sql('SELECT id, name, brand, field_provenance, field_verification FROM products ORDER BY id').all()
    assert rows[0][1:3] == ('AI jacket', 'Known brand')
    assert json.loads(rows[0][3])['name'] == 'ai_inference'
    assert json.loads(rows[0][3])['brand'] == 'legacy'
    assert json.loads(rows[0][4])['name'] == 'unverified'
    assert rows[1][1] == 'Legacy item' and json.loads(rows[1][3])['name'] == 'legacy'
    engine.dispose()


def test_inventory_migration_preserves_records_and_backfills_defaults():
    engine = create_engine('sqlite://')
    with engine.begin() as connection:
        connection.exec_driver_sql("CREATE TABLE products (id INTEGER PRIMARY KEY, name TEXT NOT NULL, category TEXT, subtype TEXT, brand TEXT, sku TEXT, color TEXT, description TEXT, condition TEXT, ai_status TEXT, assessment_status TEXT, created_at TEXT)")
        connection.exec_driver_sql("INSERT INTO products (id, name, category, created_at) VALUES (1, '', 'Electrical', '2025-04-03 12:45:00')")
        migration_path = Path(__file__).parents[1] / 'alembic' / 'versions' / '0004_inventory_fields.py'
        spec = importlib.util.spec_from_file_location('inventory_migration', migration_path)
        migration = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(migration)
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()
        row = connection.exec_driver_sql('SELECT name, category, location, completion_date, overall_condition, quantity FROM products WHERE id = 1').one()
    assert row == ('', 'Electrical', None, '2025-04-03', None, 1)
    engine.dispose()


def test_canonical_material_migration_adds_nullable_fields_without_losing_rows():
    engine = create_engine('sqlite://')
    with engine.begin() as connection:
        connection.exec_driver_sql('CREATE TABLE products (id INTEGER PRIMARY KEY, name TEXT NOT NULL)')
        connection.exec_driver_sql("INSERT INTO products (id, name) VALUES (1, 'Existing product')")
        migration_path = Path(__file__).parents[1] / 'alembic' / 'versions' / '0005_canonical_material.py'
        spec = importlib.util.spec_from_file_location('canonical_material_migration', migration_path)
        migration = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(migration)
        with Operations.context(MigrationContext.configure(connection)):
            migration.upgrade()
        columns = {column['name']: column for column in sa.inspect(connection).get_columns('products')}
        row = connection.exec_driver_sql('SELECT id, name, material, material_id, material_category FROM products').one()
    assert row == (1, 'Existing product', None, None, None)
    assert columns['material_id']['type'].__class__.__name__ == 'BIGINT'
    assert all(columns[name]['nullable'] for name in ('material', 'material_id', 'material_category', 'material_verification', 'material_confidence'))
    engine.dispose()


def test_product_images_patch_and_delete(client):
    test_client, session_factory, storage = client
    response = test_client.post('/api/products', json={
        'name': 'Jacket', 'brand': 'User brand',
        'materials': [{'name': 'Cotton', 'percentage': None}],
    })
    assert response.status_code == 201
    product_id = response.json()['id']

    response = test_client.post(f'/api/products/{product_id}/images', files=[
        ('files', ('front.png', png_bytes(), 'image/png')),
        ('files', ('back.png', png_bytes('blue'), 'image/png')),
        ('files', ('bad.png', b'not an image', 'image/png')),
    ])
    assert response.status_code == 201
    uploaded = response.json()['uploaded']
    assert len(uploaded) == 2 and len(response.json()['errors']) == 1
    assert uploaded[0]['is_primary'] is True
    assert uploaded[1]['is_primary'] is False
    assert test_client.get(uploaded[0]['url']).status_code == 200

    assert test_client.post(f"/api/products/{product_id}/images/{uploaded[1]['id']}/primary").status_code == 200
    listing_response = test_client.get('/api/products')
    assert listing_response.headers['X-Total-Count'] == '1'
    listing = listing_response.json()[0]
    assert next(image for image in listing['images'] if image['is_primary'])['id'] == uploaded[1]['id']

    response = test_client.patch(f'/api/products/{product_id}', json={'name': 'Edited jacket', 'description': 'Reviewed by user'})
    assert response.status_code == 200
    assert response.json()['name'] == 'Edited jacket'
    detail = test_client.get(f'/api/products/{product_id}').json()
    assert detail['product']['brand'] == 'User brand'
    assert len(test_client.get(f'/api/products/{product_id}/revisions').json()) == 1

    assert test_client.delete(f"/api/products/{product_id}/images/{uploaded[1]['id']}").status_code == 204
    assert test_client.get(uploaded[1]['url']).status_code == 404
    remaining = test_client.get(f'/api/products/{product_id}').json()['images']
    assert len(remaining) == 1 and remaining[0]['is_primary'] is True

    with session_factory() as db:
        db.add_all([
            AIExtraction(product_id=product_id, payload={'name': 'suggested'}, status='completed', model='test'),
            Assessment(product_id=product_id, kind='carbon', status='insufficient_data', details={}),
            Revision(product_id=product_id, snapshot={}, changed_fields=[]),
            EvidenceSource(product_id=product_id, label='test', source_type='image'),
        ])
        db.commit()
    export = test_client.get('/api/export.xlsx')
    assert export.status_code == 200 and export.content.startswith(b'PK')
    stored_name = remaining[0]['filename']
    assert test_client.delete(f'/api/products/{product_id}').status_code == 204
    assert not (storage / stored_name).exists()
    with session_factory() as db:
        assert db.get(Product, product_id) is None
        for model in (Material, ProductImage, AIExtraction, Assessment, Revision, EvidenceSource):
            assert db.scalar(select(model).where(model.product_id == product_id)) is None


def test_analysis_failure_keeps_uploaded_image_and_product(client, monkeypatch):
    test_client, _, _ = client
    created_response = test_client.post('/api/products', json={'name': '', 'brand': 'Known brand'})
    assert created_response.status_code == 201
    created = created_response.json()
    product_id = created['id']
    upload = test_client.post(f'/api/products/{product_id}/images', files={'file': ('item.png', png_bytes(), 'image/png')})
    assert upload.status_code == 201
    image_id = upload.json()['uploaded'][0]['id']

    from app.services import gemini
    monkeypatch.setattr(gemini, 'analyze_image', lambda _path, _context=None: (_ for _ in ()).throw(TimeoutError('simulated timeout')))
    failed = test_client.post(f'/api/products/{product_id}/analyze?image_id={image_id}')
    assert failed.status_code == 503
    assert test_client.get(f'/api/images/{image_id}').status_code == 200
    product = test_client.get(f'/api/products/{product_id}').json()['product']
    assert product['name'] == '' and product['brand'] == 'Known brand'
    assert product['ai_status'] == 'failed'


def test_gemini_response_returns_only_observable_product_fields(tmp_path, monkeypatch):
    from app.services import gemini

    image_path = tmp_path / 'visible.png'
    image_path.write_bytes(png_bytes())
    monkeypatch.setattr(settings, 'gemini_api_key', 'test-key')

    class FakeModels:
        def generate_content(self, **_kwargs):
            return SimpleNamespace(text='{"product_name":"Blue chair","category":"Movable Furniture and Fixed Furnishings","additional_notes":"Blue upholstered chair with four legs.","overall_condition":null,"is_intact":null,"is_clean":null,"confidence":0.91}')

    class FakeClient:
        def __init__(self, **_kwargs):
            self.models = FakeModels()

    monkeypatch.setattr(gemini.genai, 'Client', FakeClient)
    result = gemini.analyze_image(str(image_path))
    assert result['product_name'] == 'Blue chair'
    assert result['category'] == 'Movable Furniture and Fixed Furnishings'
    assert result['overall_condition'] is None and result['confidence'] == 0.91


def test_gemini_missing_fields_are_unknown_and_invalid_json_is_rejected(tmp_path, monkeypatch):
    from app.services import gemini

    image_path = tmp_path / 'visible.png'
    image_path.write_bytes(png_bytes())
    monkeypatch.setattr(settings, 'gemini_api_key', 'test-key')

    class FakeModels:
        response_text = '{}'
        def generate_content(self, **_kwargs):
            return SimpleNamespace(text=self.response_text)

    fake_models = FakeModels()
    class FakeClient:
        def __init__(self, **_kwargs):
            self.models = fake_models

    monkeypatch.setattr(gemini.genai, 'Client', FakeClient)
    missing = gemini.analyze_image(str(image_path))
    assert missing['product_name'] is None and missing['category'] is None
    fake_models.response_text = 'not-json'
    with pytest.raises(ValueError, match='invalid structured response'):
        gemini.analyze_image(str(image_path))


def test_product_can_be_created_unnamed_and_inventory_values_are_validated(client):
    test_client, _, _ = client
    created = test_client.post('/api/products', json={'product_name': '', 'location': 'Room 12', 'quantity': 2})
    assert created.status_code == 201
    assert created.json()['name'] == '' and created.json()['product_name'] == ''
    assert created.json()['location'] == 'Room 12' and created.json()['quantity'] == 2
    assert created.json()['completion_date']
    default_quantity = test_client.post('/api/products', json={'name': ''})
    assert default_quantity.status_code == 201 and default_quantity.json()['quantity'] == 1
    updated = test_client.patch(f"/api/products/{created.json()['id']}", json={'product_name': 'Identified item', 'quantity': 3, 'is_intact': True, 'overall_condition': 4})
    assert updated.status_code == 200
    assert updated.json()['name'] == 'Identified item' and updated.json()['quantity'] == 3
    assert updated.json()['is_intact'] is True and updated.json()['overall_condition'] == 4
    invalid_quantity = test_client.post('/api/products', json={'quantity': 0})
    assert invalid_quantity.status_code == 422
    invalid_category = test_client.post('/api/products', json={'category': 'Invented category'})
    assert invalid_category.status_code == 422
    assert test_client.patch(f"/api/products/{created.json()['id']}", json={'category': 'Invented category'}).status_code == 422
    invalid_condition = test_client.post('/api/products', json={'overall_condition': 6})
    assert invalid_condition.status_code == 422
