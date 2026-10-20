import io
import json
import logging
import uuid
from datetime import date, datetime
from pathlib import Path

from fastapi import APIRouter, Body, Depends, File, HTTPException, Query, Response, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from PIL import Image, UnidentifiedImageError
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.session import get_db
from app.models.product import AIExtraction, Assessment, EvidenceSource, Material, Product, ProductImage, Revision
from app.schemas.product import CATEGORY_OPTIONS, ProductIn, ProductOut, ProductUpdate, VerifyIn
from app.services.assessment import carbon_assessment, rapu_assessment
from app.services.materials import material_records, resolve_best_suggestion, resolve_material

logger = logging.getLogger(__name__)
router = APIRouter(prefix='/api')
MAX_IMAGES_PER_REQUEST = 10
ALLOWED_FORMATS = {'JPEG': ('jpg', 'image/jpeg'), 'PNG': ('png', 'image/png'), 'WEBP': ('webp', 'image/webp')}


def storage_root() -> Path:
    root = Path(settings.storage_dir).resolve()
    root.mkdir(parents=True, exist_ok=True)
    return root


def image_path(filename: str) -> Path:
    root = storage_root()
    path = (root / filename).resolve()
    if path.parent != root:
        raise HTTPException(404, 'Image not found')
    return path


def snapshot(product: Product) -> dict:
    result = {key: getattr(product, key) for key in ('id', 'name', 'category', 'location', 'completion_date', 'overall_condition', 'is_intact', 'is_clean', 'quantity', 'additional_notes', 'subtype', 'brand', 'sku', 'color', 'description', 'condition', 'material', 'material_id', 'material_category', 'material_verification', 'material_confidence', 'field_provenance', 'field_verification')}
    for key, value in result.items():
        if isinstance(value, (date, datetime)):
            result[key] = value.isoformat()
    return result


def image_data(image: ProductImage) -> dict:
    return {'id': image.id, 'filename': image.filename, 'content_type': image.content_type,
            'size_bytes': image.size_bytes, 'is_primary': image.is_primary,
            'url': f'/api/images/{image.id}'}


@router.get('/health')
def health(db: Session = Depends(get_db)):
    try:
        db.execute(select(1))
        return {'status': 'ok', 'database': 'connected'}
    except Exception as exc:
        logger.exception('Database health check failed')
        raise HTTPException(503, 'Database unavailable') from exc


@router.get('/materials')
def list_canonical_materials():
    return list(material_records())


@router.post('/products', response_model=ProductOut, status_code=201)
def create_product(data: ProductIn, db: Session = Depends(get_db)):
    values = data.model_dump(exclude={'materials'})
    mapping = resolve_material(data.material)
    values.update(material_id=mapping['material_id'] if mapping else None,
                  material_category=mapping['material_category'] if mapping else None,
                  material_verification='verified' if mapping else None,
                  material_confidence=None)
    product = Product(**values)
    product.field_provenance = {key: 'user_input' for key in ('name', 'category', 'subtype', 'brand', 'sku', 'color', 'description', 'condition')}
    product.field_verification = {key: 'unverified' for key in product.field_provenance}
    product.materials = [Material(**material.model_dump()) for material in data.materials]
    db.add(product)
    db.commit()
    db.refresh(product)
    return product


@router.get('/products')
def list_products(response: Response, q: str | None = None, status: str | None = None,
                  offset: int = Query(0, ge=0), limit: int = Query(50, ge=1, le=100),
                  db: Session = Depends(get_db)):
    stmt = select(Product)
    if q:
        stmt = stmt.where(or_(Product.name.like(f'%{q}%'), Product.brand.like(f'%{q}%'),
                              Product.sku.like(f'%{q}%'), Product.category.like(f'%{q}%')))
    if status:
        stmt = stmt.where(Product.assessment_status == status)
    total = db.scalar(select(func.count()).select_from(stmt.subquery())) or 0
    if response is not None:
        response.headers['X-Total-Count'] = str(total)
    rows = db.scalars(stmt.order_by(Product.updated_at.desc()).offset(offset).limit(limit)).all()
    result = []
    for product in rows:
        images = sorted(product.images, key=lambda item: (not item.is_primary, item.id))
        result.append({'product': ProductOut.model_validate(product).model_dump(mode='json'),
                       'materials': [{'name': m.name, 'category': m.category,
                                      'percentage': float(m.percentage) if m.percentage is not None else None,
                                      'provenance': m.provenance, 'verification': m.verification} for m in product.materials],
                       'images': [image_data(image) for image in images]})
    return result


@router.get('/products/{product_id}')
def get_product(product_id: int, db: Session = Depends(get_db)):
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(404, 'Product not found')
    images = sorted(product.images, key=lambda item: (not item.is_primary, item.id))
    materials = [{'id': m.id, 'name': m.name, 'category': m.category,
                  'percentage': float(m.percentage) if m.percentage is not None else None,
                  'provenance': m.provenance, 'verification': m.verification, 'evidence': m.evidence}
                 for m in product.materials]
    return {'product': ProductOut.model_validate(product).model_dump(mode='json'),
            'materials': materials, 'images': [image_data(image) for image in images]}


@router.patch('/products/{product_id}', response_model=ProductOut)
def update_product(product_id: int, data: ProductUpdate, db: Session = Depends(get_db)):
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(404, 'Product not found')
    changes = data.model_dump(exclude_unset=True)
    materials = changes.pop('materials', None)
    if 'material' in changes:
        mapping = resolve_material(changes['material'])
        changes.update(material=mapping['material'] if mapping else None,
                       material_id=mapping['material_id'] if mapping else None,
                       material_category=mapping['material_category'] if mapping else None,
                       material_verification='verified' if mapping else None,
                       material_confidence=None)
    if 'category' in changes and changes['category'] is not None and changes['category'] not in CATEGORY_OPTIONS and changes['category'] != product.category:
        raise HTTPException(422, 'Select one of the available product categories')
    provenance = dict(product.field_provenance or {})
    verification = dict(product.field_verification or {})
    for key in changes.keys() - {'material_id', 'material_category', 'material_verification', 'material_confidence'}:
        provenance[key] = 'user_input'
        verification[key] = 'verified'
    metadata_changed = provenance != (product.field_provenance or {}) or verification != (product.field_verification or {})
    changed = [key for key, value in changes.items() if getattr(product, key) != value]
    if metadata_changed:
        changed.append('field_verification')
    if materials is not None:
        changed.append('materials')
    if changed:
        db.add(Revision(product_id=product.id, snapshot=snapshot(product), changed_fields=changed))
        for key, value in changes.items():
            setattr(product, key, value)
        product.field_provenance = provenance
        product.field_verification = verification
        if materials is not None:
            product.materials.clear()
            product.materials.extend(Material(**item.model_dump()) for item in data.materials or [])
        db.commit()
        db.refresh(product)
    return product


async def persist_upload(product: Product, upload: UploadFile, db: Session, primary: bool) -> tuple[dict | None, str | None]:
    limit = settings.max_upload_size_mb * 1024 * 1024
    data = await upload.read(limit + 1)
    if not data or len(data) > limit:
        return None, f'{upload.filename or "Image"}: empty or exceeds {settings.max_upload_size_mb} MB'
    try:
        image = Image.open(io.BytesIO(data))
        image_format = image.format
        if image_format not in ALLOWED_FORMATS:
            raise ValueError('unsupported format')
        image.verify()
    except (UnidentifiedImageError, OSError, ValueError):
        return None, f'{upload.filename or "Image"}: upload a valid JPEG, PNG, or WebP image'
    extension, content_type = ALLOWED_FORMATS[image_format]
    filename = f'{uuid.uuid4().hex}.{extension}'
    target = image_path(filename)
    try:
        target.write_bytes(data)
        row = ProductImage(product_id=product.id, filename=filename, content_type=content_type,
                           size_bytes=len(data), is_primary=primary)
        db.add(row)
        db.flush()
        db.commit()
        db.refresh(row)
        return image_data(row), None
    except Exception:
        db.rollback()
        target.unlink(missing_ok=True)
        logger.exception('Image upload persistence failed for product %s', product.id)
        return None, f'{upload.filename or "Image"}: could not save image'


@router.post('/products/{product_id}/images', status_code=201)
async def upload_images(product_id: int, files: list[UploadFile] = File(default=[]),
                        file: UploadFile | None = File(default=None), db: Session = Depends(get_db)):
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(404, 'Product not found')
    uploads = list(files)
    if file:
        uploads.append(file)
    if not uploads:
        raise HTTPException(422, 'Choose at least one image')
    if len(uploads) > MAX_IMAGES_PER_REQUEST:
        raise HTTPException(413, f'Upload at most {MAX_IMAGES_PER_REQUEST} images at a time')
    added, errors = [], []
    for upload in uploads:
        primary = not bool(product.images) and not added
        row, error = await persist_upload(product, upload, db, primary)
        if row:
            added.append(row)
        if error:
            errors.append(error)
    if not added:
        raise HTTPException(415, {'message': 'No images were uploaded', 'errors': errors})
    return {'uploaded': added, 'errors': errors}


@router.get('/images/{image_id}')
def serve_image(image_id: int, db: Session = Depends(get_db)):
    image = db.get(ProductImage, image_id)
    if not image:
        raise HTTPException(404, 'Image not found')
    path = image_path(image.filename)
    if not path.is_file():
        raise HTTPException(404, 'Image file is missing')
    return FileResponse(path, media_type=image.content_type)


@router.post('/products/{product_id}/images/{image_id}/primary')
def select_primary_image(product_id: int, image_id: int, db: Session = Depends(get_db)):
    image = db.get(ProductImage, image_id)
    if not image or image.product_id != product_id:
        raise HTTPException(404, 'Image not found for this product')
    for candidate in db.scalars(select(ProductImage).where(ProductImage.product_id == product_id)).all():
        candidate.is_primary = candidate.id == image_id
    db.commit()
    return {'status': 'ok', 'image_id': image_id}


@router.delete('/products/{product_id}/images/{image_id}', status_code=204)
def delete_image(product_id: int, image_id: int, db: Session = Depends(get_db)):
    image = db.get(ProductImage, image_id)
    if not image or image.product_id != product_id:
        raise HTTPException(404, 'Image not found for this product')
    filename, was_primary = image.filename, image.is_primary
    db.delete(image)
    db.flush()
    if was_primary:
        next_image = db.scalar(select(ProductImage).where(ProductImage.product_id == product_id).order_by(ProductImage.id).limit(1))
        if next_image:
            next_image.is_primary = True
    db.commit()
    image_path(filename).unlink(missing_ok=True)


@router.post('/products/{product_id}/analyze')
def analyze(product_id: int, image_id: int | None = Query(default=None, gt=0),
            context: dict | None = Body(default=None), db: Session = Depends(get_db)):
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(404, 'Product not found')
    stmt = select(ProductImage).where(ProductImage.product_id == product_id)
    image = db.get(ProductImage, image_id) if image_id else db.scalar(stmt.order_by(ProductImage.is_primary.desc(), ProductImage.id).limit(1))
    if image_id and (not image or image.product_id != product_id):
        raise HTTPException(404, 'Image not found for this product')
    if not image:
        raise HTTPException(400, 'Upload an image before analysis')
    path = image_path(image.filename)
    if not path.is_file():
        raise HTTPException(404, 'Image file is missing; upload it again before analysis')
    product.ai_status = 'processing'
    db.commit()
    try:
        from app.services.gemini import analyze_image
        analysis_context = {
            'product_name': product.name or None,
            'category': product.category,
            'location': product.location,
            'additional_notes': product.additional_notes,
        }
        if context:
            for key in analysis_context:
                if key in context and isinstance(context[key], str):
                    analysis_context[key] = context[key][:2000]
        from app.schemas.product import AIProductSuggestion
        result = AIProductSuggestion.model_validate(analyze_image(str(path), analysis_context)).model_dump(mode='json')
        suggestion_name = result.pop('suggested_material', None)
        suggestion_confidence = result.pop('material_confidence', 0)
        material_candidates = result.pop('material_candidates', [])
        best_suggestion = resolve_best_suggestion(material_candidates, suggestion_name, suggestion_confidence)
        mapping = best_suggestion['mapping'] if best_suggestion else None
        best_confidence = best_suggestion['confidence'] if best_suggestion else 0
        match_score = best_suggestion['match_score'] if best_suggestion else 0
        result['material_suggestion'] = {
            'suggested_name': mapping['material'] if mapping else None,
            'matched_input': best_suggestion['suggested_name'] if best_suggestion else None,
            'material': mapping['material'] if mapping else None,
            'material_id': mapping['material_id'] if mapping else None,
            'material_category': mapping['material_category'] if mapping else None,
            'confidence': best_confidence,
            'match_score': match_score,
            'candidate_count': best_suggestion['candidate_count'] if best_suggestion else 0,
            'status': 'suggested' if best_suggestion and best_confidence >= 0.6 and match_score >= 0.78 else ('uncertain' if best_suggestion else 'unknown'),
        }
        db.add(AIExtraction(product_id=product.id, payload=result, status='completed', model=settings.gemini_model))
        product.ai_status = 'completed'
        db.commit()
        return {'status': 'completed', 'data': result, 'provenance': 'ai_inference', 'verification': 'unverified', 'image_id': image.id}
    except Exception as exc:
        logger.exception('Image analysis failed for product %s, image %s', product_id, image.id)
        db.rollback()
        product = db.get(Product, product_id)
        product.ai_status = 'failed'
        db.commit()
        raise HTTPException(503, 'Image analysis failed. Check the Gemini configuration or retry later.') from exc


@router.delete('/products/{product_id}', status_code=204)
def delete_product(product_id: int, db: Session = Depends(get_db)):
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(404, 'Product not found')
    filenames = [image.filename for image in product.images]
    try:
        db.query(AIExtraction).filter(AIExtraction.product_id == product_id).delete(synchronize_session=False)
        db.query(Assessment).filter(Assessment.product_id == product_id).delete(synchronize_session=False)
        db.query(Revision).filter(Revision.product_id == product_id).delete(synchronize_session=False)
        db.query(EvidenceSource).filter(EvidenceSource.product_id == product_id).delete(synchronize_session=False)
        db.delete(product)
        db.commit()
    except Exception as exc:
        db.rollback()
        logger.exception('Product deletion failed for product %s', product_id)
        raise HTTPException(500, 'Could not delete product') from exc
    for filename in filenames:
        image_path(filename).unlink(missing_ok=True)


@router.post('/products/{product_id}/assessments')
def assess(product_id: int, db: Session = Depends(get_db)):
    product = db.get(Product, product_id)
    if not product:
        raise HTTPException(404, 'Product not found')
    output = []
    for result in (carbon_assessment(product), rapu_assessment(product)):
        db.add(Assessment(product_id=product.id, kind=result['kind'], status=result['status'], result=result['result'],
                          unit=result['unit'], methodology=result['methodology'], details=result['details']))
        output.append(result)
    db.commit()
    return output


@router.get('/products/{product_id}/revisions')
def revisions(product_id: int, db: Session = Depends(get_db)):
    if not db.get(Product, product_id):
        raise HTTPException(404, 'Product not found')
    rows = db.scalars(select(Revision).where(Revision.product_id == product_id).order_by(Revision.created_at.desc())).all()
    return [{'id': row.id, 'snapshot': row.snapshot, 'changed_fields': row.changed_fields,
             'created_at': row.created_at.isoformat()} for row in rows]


@router.post('/products/{product_id}/materials/{material_id}/verify')
def verify(product_id: int, material_id: int, data: VerifyIn, db: Session = Depends(get_db)):
    material = db.get(Material, material_id)
    if not material or material.product_id != product_id:
        raise HTTPException(404, 'Material not found')
    material.verification = data.verification.value
    material.evidence = data.reviewer_note
    db.add(Revision(product_id=product_id, snapshot={'material_id': material.id, 'verification': material.verification,
                                                      'note': material.evidence}, changed_fields=['material.verification']))
    db.commit()
    return {'status': material.verification}


def safe_cell(value):
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + value
    return value if value is not None else 'Unknown'


@router.get('/export.xlsx')
def export_products(db: Session = Depends(get_db)):
    products = db.scalars(select(Product).order_by(Product.id)).all()
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Product summary'
    summary.append(['ID', 'Name', 'Category', 'Subtype', 'Brand', 'SKU', 'Color', 'Condition', 'Material', 'Material ID', 'Material Category', 'Material verification', 'AI status', 'Assessment status'])
    materials = workbook.create_sheet('Material composition')
    materials.append(['Product ID', 'Product', 'Material', 'Category', 'Percentage by weight', 'Provenance', 'Verification', 'Evidence'])
    sustainability = workbook.create_sheet('Sustainability')
    sustainability.append(['Product ID', 'Product', 'Assessment', 'Status', 'Result', 'Unit', 'Methodology', 'Details'])
    evidence = workbook.create_sheet('Evidence')
    evidence.append(['Product ID', 'Product', 'Image', 'Reference', 'Verification'])
    for product in products:
        summary.append([safe_cell(getattr(product, key)) for key in ('id', 'name', 'category', 'subtype', 'brand', 'sku', 'color', 'condition', 'material', 'material_id', 'material_category', 'material_verification', 'ai_status', 'assessment_status')])
        for material in product.materials:
            materials.append([safe_cell(value) for value in [product.id, product.name, material.name, material.category, material.percentage, material.provenance, material.verification, material.evidence]])
        for assessment in db.scalars(select(Assessment).where(Assessment.product_id == product.id)):
            sustainability.append([safe_cell(value) for value in [product.id, product.name, assessment.kind, assessment.status, assessment.result, assessment.unit, assessment.methodology, json.dumps(assessment.details)]])
        for image in product.images:
            evidence.append([safe_cell(value) for value in [product.id, product.name, image.filename, None, 'unverified']])
    for sheet in workbook.worksheets:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(42, max(12, max((len(str(cell.value or '')) for cell in column), default=10) + 2))
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(output, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                             headers={'Content-Disposition': 'attachment; filename="rakki-products.xlsx"'})
